from datetime import datetime, time
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.appointments.models import Cita
from apps.clients.models import Cliente
from apps.common.custom_time_fields import MONTH_NUMBER_TO_NAME
from apps.payments.choices import EstadoPago
from apps.payments.models import Pago
from apps.profiles.models import Perfil
from tests.base import NegocioTestCase

User = get_user_model()


class CompletadosTests(NegocioTestCase):
    esquema = "completados"
    dominio = "completados.localhost"

    def setUp(self):
        super().setUp()
        encargada = User.objects.create_user(
            username="ana", email="ana@correo.cl", password="x"
        )
        Perfil.objects.filter(user=encargada).update(
            rol=Perfil.Rol.ENCARGADO, debe_cambiar_clave=False
        )
        self.navegador = Client(HTTP_HOST=self.dominio)
        self.navegador.force_login(encargada)
        self.cliente = Cliente.objects.create(nombre="Paula", apellido="Herrera")
        hoy = timezone.localdate()
        self.momento = timezone.make_aware(
            datetime.combine(hoy.replace(day=1), time(10, 0))
        )
        self.mes = f"{MONTH_NUMBER_TO_NAME[hoy.month]} {hoy.year}"
        self.pago("Atendida", Cita.EstadoChoices.COMPLETADA, "20000")
        self.pago("Cancelada", Cita.EstadoChoices.CANCELADA, "5000")

    def pago(self, nombre, estado_cita, total):
        cita = Cita.objects.create(
            cliente=self.cliente,
            fecha_agenda=self.momento.date(),
            hora_agenda=self.momento.time(),
            estado=estado_cita,
        )
        return Pago.objects.create(
            cita=cita,
            monto_total_cita=Decimal(total),
            estado_pago=EstadoPago.COMPLETADO,
            fecha_pago_completado=self.momento,
            cliente_nombre=nombre,
            fecha_cita=self.momento,
        )

    def test_la_lista_solo_trae_citas_completadas(self):
        datos = self.navegador.get(
            reverse("payments_list"),
            {"draw": 1, "start": 0, "length": 10, "months": self.mes},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        ).json()
        self.assertEqual(
            [fila["cliente_nombre"] for fila in datos["data"]], ["Atendida"]
        )
        self.assertEqual(
            (datos["monto_total_cita"], datos["monto_recibido_mes"]),
            ("$20.000", "$20.000"),
        )

    def test_el_grafico_semanal_no_suma_la_cancelada(self):
        grafico = self.navegador.get(
            reverse("payments_weekly_income_ajax"),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        ).json()
        self.assertEqual(sum(grafico["datasets"][0]["data"]), 20000)

    def test_el_excel_no_trae_la_cancelada(self):
        hoja = load_workbook(
            BytesIO(
                self.navegador.get(
                    reverse("payments_export"), {"months": self.mes}
                ).content
            )
        ).active
        clientes = [
            fila[1] for fila in hoja.iter_rows(min_row=2, values_only=True) if fila[1]
        ]
        self.assertEqual(clientes, ["Atendida"])

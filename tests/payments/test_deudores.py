from datetime import date, time
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.appointments.models import Cita
from apps.clients.models import Cliente
from apps.payments.choices import EstadoDetallePago, EstadoPago, MetodoPago, TipoDetallePago
from apps.payments.models import DetallePago, Pago
from apps.profiles.models import Perfil
from tests.base import NegocioTestCase

User = get_user_model()
Tipo = TipoDetallePago
Estado = EstadoDetallePago


class DeudoresConAnticiposTests(NegocioTestCase):
    esquema = "deudores_anticipos"
    dominio = "deudores-anticipos.localhost"

    def setUp(self):
        super().setUp()
        encargada = User.objects.create_user(username="ana", email="ana@correo.cl", password="x")
        Perfil.objects.filter(user=encargada).update(rol=Perfil.Rol.ENCARGADO, debe_cambiar_clave=False)
        self.navegador = Client(HTTP_HOST=self.dominio)
        self.navegador.force_login(encargada)
        cita = Cita.objects.create(
            cliente=Cliente.objects.create(nombre="Paula", apellido="Herrera"),
            fecha_agenda=date(2026, 10, 25),
            hora_agenda=time(10, 0),
            estado=Cita.EstadoChoices.COMPLETADA,
        )
        self.pago = Pago.objects.create(
            cita=cita,
            monto_total_cita=Decimal("38000"),
            estado_pago=EstadoPago.PENDIENTE,
            cliente_nombre="Paula Herrera",
            fecha_cita=timezone.now(),
        )
        self.fila("3000", Tipo.RESERVA, Estado.VIGENTE, MetodoPago.TRANSFERENCIA)
        self.fila("5000", Tipo.ANTICIPO, Estado.VIGENTE)
        self.fila("10000", Tipo.ABONO, Estado.VIGENTE)
        self.fila("8000", Tipo.ANTICIPO, Estado.RETENIDO_DEFINITIVO)
        devuelto = self.fila("4000", Tipo.ANTICIPO, Estado.DEVUELTO)
        self.fila("4000", Tipo.DEVOLUCION, Estado.VIGENTE, devuelve=devuelto)

    def fila(self, monto, tipo, estado, metodo=MetodoPago.EFECTIVO, devuelve=None):
        return DetallePago.objects.create(
            pago=self.pago,
            fecha_pago=timezone.now(),
            monto_pago=Decimal(monto),
            metodo_pago=metodo,
            tipo_detalle=tipo,
            estado_detalle=estado,
            detalle_devuelto=devuelve,
        )

    def ajax(self, nombre, datos=None, **kwargs):
        respuesta = self.navegador.get(
            reverse(nombre, kwargs=kwargs or None), datos or {}, HTTP_X_REQUESTED_WITH="XMLHttpRequest"
        )
        self.assertEqual(respuesta.status_code, 200)
        return respuesta

    def test_la_lista_muestra_la_cuenta_de_la_cita_sin_lo_retenido(self):
        datos = self.ajax("debtors_list", {"draw": 1, "start": 0, "length": 10}).json()
        fila = datos["data"][0]
        self.assertEqual(
            (fila["total_cita_formatted"], fila["total_abonado_formatted"], fila["saldo_pendiente_formatted"]),
            ("$ 30,000", "$ 18,000", "$ 12,000"),
        )
        self.assertEqual((datos["total_pendiente"], datos["total_abonado"]), ("$ 12,000", "$ 18,000"))

    def test_la_tarjeta_total_pendiente_resta_lo_abonado(self):
        otra = Pago.objects.create(
            cita=Cita.objects.create(
                cliente=self.pago.cita.cliente,
                fecha_agenda=date(2026, 10, 26),
                hora_agenda=time(10, 0),
                estado=Cita.EstadoChoices.COMPLETADA,
            ),
            monto_total_cita=Decimal("20000"),
            estado_pago=EstadoPago.PENDIENTE,
            cliente_nombre="Paula Herrera",
            fecha_cita=timezone.now(),
        )
        DetallePago.objects.create(
            pago=otra, fecha_pago=timezone.now(), monto_pago=Decimal("5000"), metodo_pago=MetodoPago.EFECTIVO
        )
        datos = self.ajax("debtors_list", {"draw": 1, "start": 0, "length": 10}).json()
        self.assertEqual((datos["total_pendiente"], datos["total_abonado"]), ("$ 27,000", "$ 23,000"))

    def test_el_detalle_de_deuda_cuadra_con_la_lista(self):
        contexto = self.ajax("debt_detail_modal", pk=self.pago.pk).context
        self.assertEqual(
            (contexto["total_cita"], contexto["total_pagado"], contexto["total_pendiente"]),
            ("$ 30,000", "$ 18,000", "$ 12,000"),
        )

    def test_los_pagos_del_detalle_solo_son_reserva_anticipo_y_abono(self):
        datos = self.ajax("payment_detail_list", {"draw": 1}, pk=self.pago.pk).json()
        filas = sorted((fila["tipo_detalle_display"], fila["monto_pago_formatted"]) for fila in datos["data"])
        self.assertEqual(filas, [("Abono", "$ 10,000"), ("Anticipo", "$ 5,000"), ("Reserva", "$ 3,000")])

    def test_las_filas_antiguas_sin_tipo_salen_como_abono(self):
        DetallePago.objects.create(
            pago=self.pago, fecha_pago=timezone.now(), monto_pago=Decimal("1000"), metodo_pago=MetodoPago.EFECTIVO
        )
        datos = self.ajax("payment_detail_list", {"draw": 1}, pk=self.pago.pk).json()
        self.assertIn(("Abono", "$ 1,000"), [(f["tipo_detalle_display"], f["monto_pago_formatted"]) for f in datos["data"]])

    def test_abonar_deuda_muestra_las_mismas_cifras_y_pagos(self):
        contexto = self.ajax("add_payment_modal", pk=self.pago.pk).context
        self.assertEqual((contexto["total"], contexto["pagado"], contexto["saldo"]), ("$30.000", "$18.000", "$12.000"))
        self.assertEqual(contexto["payment_config"], {"total": 30000, "paid": 18000, "balance": 12000})
        self.assertEqual(
            sorted((abono["kind"], abono["amount"]) for abono in contexto["abonos"]),
            [("Abono", "$10.000"), ("Anticipo", "$5.000"), ("Reserva", "$3.000")],
        )

    def test_pagar_el_saldo_completa_el_pago(self):
        respuesta = self.navegador.post(
            reverse("add_payment_modal", args=[self.pago.pk]),
            {
                "full_payment": "on",
                "payment_method": MetodoPago.EFECTIVO,
                "payment_date": timezone.localdate().isoformat(),
                "payment_time": "10:00",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(respuesta.status_code, 200)
        self.pago.refresh_from_db()
        self.assertEqual(self.pago.estado_pago, EstadoPago.COMPLETADO)
        self.assertEqual(DetallePago.objects.filter(tipo_detalle=Tipo.ABONO).latest("pk").monto_pago, Decimal("12000"))

    def test_el_excel_trae_las_mismas_cifras(self):
        respuesta = self.navegador.get(reverse("debtors_export"))
        hoja = load_workbook(BytesIO(respuesta.content)).active
        self.assertEqual([celda.value for celda in hoja[1]][2:], ["Total de la cita", "Total abonado", "Total pendiente"])
        self.assertEqual(list(next(hoja.iter_rows(min_row=2, values_only=True)))[2:], ["$ 30,000", "$ 18,000", "$ 12,000"])

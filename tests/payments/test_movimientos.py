from datetime import date, time
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.appointments.models import Cita
from apps.clients.models import Cliente
from apps.payments.choices import EstadoDetallePago, MetodoPago, TipoDetallePago, TipoMovimiento
from apps.payments.models import DetallePago, Pago
from apps.profiles.models import Perfil
from tests.base import NegocioTestCase

User = get_user_model()
Tipo = TipoDetallePago
Estado = EstadoDetallePago
Movimiento = TipoMovimiento


class MovimientosTests(NegocioTestCase):
    esquema = "movimientos"
    dominio = "movimientos.localhost"

    def setUp(self):
        super().setUp()
        encargada = User.objects.create_user(username="ana", email="ana@correo.cl", password="x")
        Perfil.objects.filter(user=encargada).update(rol=Perfil.Rol.ENCARGADO, debe_cambiar_clave=False)
        self.navegador = Client(HTTP_HOST=self.dominio)
        self.navegador.force_login(encargada)
        cita = Cita.objects.create(
            cliente=Cliente.objects.create(nombre="Paula", apellido="Herrera"),
            fecha_agenda=date(2026, 10, 20),
            hora_agenda=time(10, 0),
            estado=Cita.EstadoChoices.CANCELADA,
        )
        self.pago = Pago.objects.create(
            cita=cita,
            monto_total_cita=Decimal("6000"),
            cliente_nombre="Paula Herrera",
            fecha_cita=timezone.now(),
        )
        self.antigua = self.fila("3000")
        self.abono = self.fila("2000", Tipo.ABONO, Estado.VIGENTE, metodo=MetodoPago.TARJETA)
        self.retenido = self.fila("6000", Tipo.ANTICIPO, Estado.RETENIDO)
        devuelto = self.fila("4000", Tipo.ANTICIPO, Estado.DEVUELTO, metodo=MetodoPago.TRANSFERENCIA)
        self.reembolso = self.fila(
            "4000", Tipo.DEVOLUCION, Estado.VIGENTE, metodo=MetodoPago.TRANSFERENCIA, devuelve=devuelto
        )

    def fila(self, monto, tipo=None, estado=None, metodo=MetodoPago.EFECTIVO, devuelve=None):
        return DetallePago.objects.create(
            pago=self.pago,
            fecha_pago=timezone.now(),
            monto_pago=Decimal(monto),
            metodo_pago=metodo,
            tipo_detalle=tipo,
            estado_detalle=estado,
            detalle_devuelto=devuelve,
        )

    def listar(self, **filtros):
        respuesta = self.navegador.get(
            reverse("incomes_list"),
            {"draw": 1, "start": 0, "length": 25, **filtros},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(respuesta.status_code, 200)
        return respuesta.json()

    def filas(self, **filtros):
        return {fila["monto_movimiento_formatted"]: fila for fila in self.listar(**filtros)["data"]}

    def test_la_pantalla_se_llama_movimientos_y_ofrece_el_filtro_de_movimiento(self):
        respuesta = self.navegador.get(reverse("incomes"))
        self.assertContains(respuesta, "Movimientos")
        self.assertContains(respuesta, "Total de egresos")
        self.assertContains(respuesta, "Tipo de movimiento")
        self.assertContains(respuesta, "Ingresos por método de pago")
        self.assertNotContains(respuesta, "Monto recibido")
        self.assertNotContains(respuesta, "Cantidad")
        opciones = respuesta.context["filter_form"].fields["movement"].widget.choices
        self.assertEqual([etiqueta for _, etiqueta in opciones], ["Todos", "Ingresos", "Egresos"])
        menu = self.navegador.get(reverse("calendar"))
        self.assertContains(menu, "Movimientos")
        self.assertNotContains(menu, "Ingresos")

    def test_cada_fila_dice_si_es_ingreso_o_egreso_y_el_egreso_va_en_negativo(self):
        filas = self.filas()
        self.assertEqual(len(filas), 5)
        egreso = filas["$ -4,000"]
        self.assertEqual((egreso["tipo_movimiento_display"], egreso["es_egreso"]), ("Egreso", True))
        for monto in ("$ 3,000", "$ 2,000", "$ 6,000", "$ 4,000"):
            with self.subTest(monto=monto):
                self.assertEqual((filas[monto]["tipo_movimiento_display"], filas[monto]["es_egreso"]), ("Ingreso", False))

    def test_los_totales_separan_ingresos_egresos_y_neto(self):
        datos = self.listar()
        self.assertEqual(
            (datos["total_ingresos"], datos["total_egresos"], datos["total_neto"]),
            ("$ 15,000", "$ 4,000", "$ 11,000"),
        )
        self.assertNotIn("cantidad_movimientos", datos)

    def test_filtra_ingresos_o_egresos(self):
        self.assertEqual(sorted(self.filas(movement=Movimiento.INGRESO)), ["$ 2,000", "$ 3,000", "$ 4,000", "$ 6,000"])
        datos = self.listar(movement=Movimiento.EGRESO)
        self.assertEqual([fila["monto_movimiento_formatted"] for fila in datos["data"]], ["$ -4,000"])
        self.assertEqual((datos["total_ingresos"], datos["total_egresos"], datos["total_neto"]), ("$ 0", "$ 4,000", "$ -4,000"))
        self.assertEqual(len(self.filas(movement="OTRO")), 5)

    def test_combina_movimiento_con_metodo(self):
        filas = self.filas(movement=Movimiento.INGRESO, payment_method=MetodoPago.TRANSFERENCIA)
        self.assertEqual(list(filas), ["$ 4,000"])

    def test_ordena_por_monto_con_signo(self):
        datos = self.listar(**{"order[0][column]": 6, "order[0][dir]": "asc"})
        self.assertEqual(datos["data"][0]["monto_movimiento_formatted"], "$ -4,000")

    def test_las_filas_con_la_misma_fecha_no_se_repiten_entre_paginas(self):
        momento = timezone.now()
        DetallePago.objects.update(fecha_pago=momento)
        orden = {"order[0][column]": 0, "order[0][dir]": "desc"}
        vistas = [
            self.listar(start=inicio, length=2, **orden)["data"] for inicio in (0, 2, 4)
        ]
        montos = [fila["monto_movimiento_formatted"] for pagina in vistas for fila in pagina]
        self.assertEqual(montos, ["$ -4,000", "$ 4,000", "$ 6,000", "$ 2,000", "$ 3,000"])

    def grafico(self, **filtros):
        respuesta = self.navegador.get(
            reverse("incomes_by_method_ajax"), filtros, HTTP_X_REQUESTED_WITH="XMLHttpRequest"
        )
        return respuesta.json()

    def test_el_grafico_solo_muestra_ingresos_por_metodo(self):
        grafico = self.grafico()
        dataset = grafico["datasets"][0]
        datos = dict(zip(dataset["keys"], dataset["data"]))
        self.assertEqual(dataset["label"], "Recibido")
        self.assertEqual((datos["efectivo"], datos["tarjeta"], datos["transferencia"]), (9000, 2000, 4000))

    def test_con_egresos_el_grafico_queda_vacio(self):
        self.assertTrue(self.grafico(movement=Movimiento.EGRESO)["meta"]["empty"])

    def test_el_excel_trae_tipo_de_movimiento_y_montos_con_signo(self):
        respuesta = self.navegador.get(reverse("incomes_export"), {"movement": ""})
        self.assertIn("movimientos_", respuesta["Content-Disposition"])
        hoja = load_workbook(BytesIO(respuesta.content)).active
        self.assertEqual(hoja.title, "Movimientos")
        encabezados = [celda.value for celda in hoja[1]]
        self.assertEqual(encabezados[-2:], ["Tipo de movimiento", "Monto"])
        filas = {fila[-1]: fila[-2] for fila in hoja.iter_rows(min_row=2, values_only=True) if fila[-1]}
        self.assertEqual(filas["$ -4,000"], "Egreso")
        self.assertEqual(filas["$ 6,000"], "Ingreso")
        self.assertEqual(filas["$ 3,000"], "Ingreso")

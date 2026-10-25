import io
from types import SimpleNamespace

from django.db import connection
from django.test import RequestFactory
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from apps.clients.imports.validators import ClientAsyncImportValidator
from apps.clients.models import Cliente
from apps.clients.views.imports import ClientExampleExportView
from tests.base import NegocioTestCase

ENCABEZADO = (
    "Nombre,Apellido,Teléfono,Email,"
    "País,Tipo de identificador nacional,Identificador nacional,Estado,Notas"
)


def validar(*filas):
    contenido = "\n".join([ENCABEZADO, *filas]) + "\n"
    validador = ClientAsyncImportValidator(
        SimpleNamespace(datos_entrada={"contenido": contenido})
    )
    return validador.validate()


def identificador(fila):
    return (
        fila["pais"],
        fila["tipo_identificador_nacional"],
        fila["identificador_nacional"],
    )


def valores(hoja):
    return [list(fila) for fila in hoja.iter_rows(values_only=True)]


class ImportarClientesTests(NegocioTestCase):
    esquema = "importar_clientes"

    def test_las_columnas_siguen_el_orden_de_la_plantilla(self):
        self.assertEqual(
            ClientAsyncImportValidator.get_headers(), ENCABEZADO.split(",")
        )
        for fila in ClientExampleExportView.example_rows:
            self.assertEqual(len(fila), len(ENCABEZADO.split(",")))
        self.assertTrue(
            validar(",".join(ClientExampleExportView.example_rows[0])).is_ok()
        )

    def test_la_plantilla_trae_las_hojas_de_ayuda(self):
        respuesta = ClientExampleExportView().get(RequestFactory().get("/"))
        libro = load_workbook(io.BytesIO(respuesta.content))
        self.assertEqual(
            libro.sheetnames, ["Plantilla", "Países", "Tipos de identificador"]
        )
        self.assertEqual(valores(libro["Plantilla"])[0], ENCABEZADO.split(","))
        paises = valores(libro["Países"])
        self.assertEqual(paises[0], ["Código", "País"])
        self.assertIn(["CL", "Chile"], paises)
        self.assertIn(["VE", "Venezuela"], paises)
        self.assertEqual(len(paises), 10)
        self.assertEqual(
            valores(libro["Tipos de identificador"]),
            [
                ["Código", "Tipo"],
                ["cedula", "Cédula"],
                ["curp", "CURP"],
                ["dni", "DNI"],
                ["pasaporte", "Pasaporte"],
                ["rut", "RUT"],
            ],
        )

    def test_el_identificador_es_opcional_y_el_pais_puede_ir_solo(self):
        resultado = validar(
            "Ana,Rojas,,,,,,activo,",
            "Lucía,Soto,,,ve,,,activo,",
        )
        self.assertTrue(resultado.is_ok())
        self.assertEqual(
            [identificador(fila) for fila in resultado.ok_value],
            [(None, None, None), ("VE", None, None)],
        )

    def test_acepta_el_tipo_por_codigo_o_nombre_y_normaliza_el_numero(self):
        resultado = validar(
            "Ana,Rojas,,,cl,RUT,12.345.678-5,activo,",
            "Lucía,Soto,,,VE,Cédula de identidad,v-12.345.678,activo,",
            "Carla,Díaz,,,UY,Cédula,1.234.567-2,activo,",
            "Rosa,Paz,,,PE,pasaporte,ab123456,activo,",
        )
        self.assertTrue(resultado.is_ok())
        self.assertEqual(
            [identificador(fila) for fila in resultado.ok_value],
            [
                ("CL", "rut", "12345678-5"),
                ("VE", "cedula", "V-12345678"),
                ("UY", "cedula", "1234567-2"),
                ("PE", "pasaporte", "AB123456"),
            ],
        )

    def test_aplica_las_mismas_reglas_que_el_formulario(self):
        casos = {
            "Ana,Rojas,,,CL,rut,,activo,": "Escribe el identificador nacional del tipo elegido.",
            "Ana,Rojas,,,CL,,12345678-5,activo,": "Elige el tipo de identificador nacional.",
            "Ana,Rojas,,,,rut,12345678-5,activo,": "Elige el país del identificador nacional.",
            "Ana,Rojas,,,ES,dni,12345678,activo,": "El país debe ser un código de dos letras",
            "Ana,Rojas,,,CL,curp,BOXW310820HNERXN09,activo,": "El tipo de identificador no corresponde a este país.",
            "Ana,Rojas,,,CL,rut,12.345.678-9,activo,": "dígito verificador",
        }
        for fila, mensaje in casos.items():
            with self.subTest(fila=fila):
                resultado = validar(fila)
                self.assertTrue(resultado.is_err())
                self.assertTrue(
                    any(mensaje in error["message"] for error in resultado.err_value)
                )

    def test_no_repite_datos_de_clientes_existentes_ni_dentro_del_archivo(self):
        Cliente.objects.create(
            nombre="Otra",
            email="ana@correo.cl",
            telefono="+56912345678",
            pais="CL",
            tipo_identificador_nacional="rut",
            identificador_nacional="12345678-5",
        )
        contra_la_base = validar(
            "Ana,Rojas,+56912345678,ANA@correo.cl,CL,rut,12.345.678-5,activo,"
        )
        mensajes = {error["message"] for error in contra_la_base.err_value}
        esperados = {
            "Ese correo ya lo usa otro cliente.",
            "Ese teléfono ya lo usa otro cliente.",
            "Ese identificador nacional ya lo tiene otro cliente.",
        }
        self.assertTrue(esperados <= mensajes)
        dentro_del_archivo = validar(
            "Ana,Rojas,+56987654321,ana2@correo.cl,VE,cedula,V-12345678,activo,",
            "Lucía,Soto,+56987654321,ANA2@correo.cl,ve,cedula,v-12.345.678,activo,",
        )
        mensajes = sorted(
            (error["row"], error["message"]) for error in dentro_del_archivo.err_value
        )
        self.assertEqual(
            mensajes,
            [
                (3, "Ese correo se repite en otra fila del archivo."),
                (3, "Ese identificador nacional se repite en otra fila del archivo."),
                (3, "Ese teléfono se repite en otra fila del archivo."),
            ],
        )

    def test_las_consultas_no_crecen_con_la_cantidad_de_filas(self):
        def consultas(cantidad):
            filas = [
                f"Importado,X,+5698{i:07d},imp{i}@correo.cl,,,,activo,"
                for i in range(cantidad)
            ]
            with CaptureQueriesContext(connection) as contexto:
                self.assertTrue(validar(*filas).is_ok())
            return len(contexto.captured_queries)

        self.assertEqual(consultas(5), consultas(40))

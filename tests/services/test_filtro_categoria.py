from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from apps.profiles.models import Perfil
from apps.services.models import Categoria, Servicio
from tests.base import NegocioTestCase

User = get_user_model()


class FiltroCategoriaServiciosTests(NegocioTestCase):
    esquema = "filtro_categoria"
    dominio = "filtro-categoria.localhost"

    def setUp(self):
        super().setUp()
        encargada = User.objects.create_user(username="ana", email="ana@correo.cl", password="x")
        Perfil.objects.filter(user=encargada).update(rol=Perfil.Rol.ENCARGADO, debe_cambiar_clave=False)
        self.navegador = Client(HTTP_HOST=self.dominio)
        self.navegador.force_login(encargada)
        self.unas = Categoria.objects.create(nombre="Uñas")
        self.pestanas = Categoria.objects.create(nombre="Pestañas")
        Servicio.objects.create(nombre="Acrílico", precio=Decimal("15000"), categoria=self.unas)
        Servicio.objects.create(nombre="Lifting", precio=Decimal("20000"), categoria=self.pestanas)
        Servicio.objects.create(nombre="Masaje", precio=Decimal("10000"))
        Servicio.objects.create(
            nombre="Retiro", precio=Decimal("5000"), estado=Servicio.EstadoChoices.INACTIVO
        )

    def listar(self, **filtros):
        respuesta = self.navegador.get(
            reverse("service_list"),
            {"draw": 1, "start": 0, "length": 10, "status": "all", "category": "all", **filtros},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(respuesta.status_code, 200)
        return sorted(fila["nombre"] for fila in respuesta.json()["data"])

    def test_todas_trae_todos_los_servicios(self):
        self.assertEqual(self.listar(), ["Acrílico", "Lifting", "Masaje", "Retiro"])

    def test_filtra_por_una_categoria(self):
        self.assertEqual(self.listar(category=str(self.unas.pk)), ["Acrílico"])

    def test_sin_definir_trae_los_servicios_sin_categoria(self):
        self.assertEqual(self.listar(category="undefined"), ["Masaje", "Retiro"])

    def test_combina_categoria_con_estado(self):
        self.assertEqual(self.listar(category="undefined", status="activo"), ["Masaje"])

    def test_la_pagina_ofrece_todas_sin_definir_y_cada_categoria(self):
        respuesta = self.navegador.get(reverse("services"))
        opciones = respuesta.context["filter_form"].fields["category"].choices
        self.assertEqual(
            [etiqueta for _, etiqueta in opciones],
            ["Todas", "Sin definir", "Pestañas", "Uñas"],
        )
        self.assertContains(respuesta, "form-select-2--custom")

    def test_exporta_respetando_la_categoria(self):
        respuesta = self.navegador.get(
            reverse("service_export"), {"status": "all", "category": "undefined"}
        )
        hoja = load_workbook(BytesIO(respuesta.content)).active
        nombres = [fila[0] for fila in hoja.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(filter(None, nombres)), ["Masaje", "Retiro"])
